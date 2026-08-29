#!/usr/bin/env python3
"""Parse KIB B/C/D Excel: nama barang, tipe, lokasi, tahun anggaran."""
from __future__ import annotations

import gzip
import json
import re
from pathlib import Path

import openpyxl

ROOT = Path("/workspace")
ATT = ROOT / "attachments"

# Koordinat tidak diisi dari perkiraan. Hanya nama lokasi dari KIB;
# titik GPS diisi operator lewat kamera geotag.

PRODUCT_RE = re.compile(
    r"pju|lpju|apj|lampu|led|watt|armature|siklon|tiang|oktagonal|solar|"
    r"rambu|apill|traffic|marka|guardrail|warning light|stang ornamen",
    re.I,
)
PROCUREMENT_RE = re.compile(
    r"^(pengadaan|belanja|pemasangan|id awal|hasil clustering|di tambah|keg[\.\s]|---)",
    re.I,
)
JUNK_LINE = re.compile(
    r"^(id awal|gambar\s*:|skpd:|---|\s*-+\s*$|kel/desa\.|kec\.|kab\.|kota |rt/rw)",
    re.I,
)
JUNK_NAME = re.compile(
    r"bast|mutasi:|id awal|kp/komp|gambar\s*:|skpd:|inv[-.]?\d",
    re.I,
)
OUTDOOR_B = re.compile(
    r"rambu|traffic light|lampu natrium|lampu untuk menerangi|apill|pju|"
    r"marka jalan|guardrail|delineator|warning light|tiang lampu|lampu jalan|"
    r"solar cell|lampu pijar|lampu led",
    re.I,
)
VEHICLE_B = re.compile(
    r"jeep|bus|minibus|mini bus|sepeda motor|sepeda|mobil|station wagon|"
    r"pickup|pick up|truck|truk|ambulance|loader|sedan|mpv|hiace|pajero|"
    r"avanza|rush|fortuner|hilux|ertiga",
    re.I,
)


def clean(v) -> str:
    if v is None:
        return ""
    s = str(v).replace("\xa0", " ").replace("&", "&").replace("\r", " ")
    s = re.sub(r"[ \t]+", " ", s)
    s = re.sub(r"\n+", "\n", s).strip()
    return s


def first_line(v) -> str:
    for line in clean(v).split("\n"):
        line = line.strip(" -/")
        if not line or line in {"-", "–"} or JUNK_LINE.match(line):
            continue
        line = re.split(r"\s*RT/RW", line, maxsplit=1)[0].strip(" -/")
        if line:
            return line[:160]
    return ""


def is_main_no(v) -> bool:
    if v is None or isinstance(v, bool):
        return False
    if isinstance(v, (int, float)):
        return int(v) == v and int(v) >= 1
    return str(v).strip().isdigit()


def parse_ids(raw) -> tuple[str, str]:
    s = clean(raw).replace("\n", "/")
    parts = [p.strip() for p in s.split("/") if p.strip()]
    kode = re.sub(r"\s+", "", parts[0]) if parts else ""
    idb = re.sub(r"\D", "", parts[1]) if len(parts) > 1 else ""
    if not idb and len(parts) > 2:
        idb = re.sub(r"\D", "", parts[2])
    return kode, idb


def parse_year(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        y = int(v)
        return y if 1980 <= y <= 2026 else None
    m = re.search(r"(19|20)\d{2}", str(v))
    return int(m.group()) if m else None


def parse_qty(v) -> int:
    if v is None:
        return 1
    if isinstance(v, (int, float)):
        n = int(v)
        return n if n > 0 else 1
    m = re.search(r"\d+", str(v).replace(".", "").replace(",", ""))
    if not m:
        return 1
    n = int(m.group())
    return n if 0 < n < 100000 else 1


def parse_cond(*vals) -> str:
    t = " ".join(clean(v).lower() for v in vals)
    if "rusak berat" in t or re.search(r"\brb\b", t):
        return "rusak_berat"
    if "rusak ringan" in t or re.search(r"\bkb\b", t):
        return "rusak_ringan"
    if "hilang" in t:
        return "hilang"
    return "baik"


def dash_none(s: str) -> str:
    t = (s or "").strip().strip("-–.")
    return t if t and t not in {"-", "/"} else ""


def format_plate(s: str) -> str:
    raw = dash_none(s).replace(".", " ")
    m = re.match(r"^([A-Za-z]{1,2})\s*(\d{1,4})\s*([A-Za-z]{0,3})$", raw)
    if m:
        return " ".join(p.upper() for p in m.groups() if p)
    return raw.upper() if raw else ""


def parse_reg(v) -> str:
    s = clean(v).split("\n")[0]
    s = re.sub(r"[\s/]+", "", s)
    return s[:24]


def kab_from(text: str) -> str:
    t = clean(text).upper()
    if "TANGERANG SELATAN" in t:
        return "KOTA TANGERANG SELATAN"
    if "KOTA TANGERANG" in t:
        return "KOTA TANGERANG"
    if "KAB" in t and "TANGERANG" in t:
        return "KAB. TANGERANG"
    if "KOTA SERANG" in t:
        return "KOTA SERANG"
    if "SERANG" in t:
        return "KAB. SERANG"
    if "CILEGON" in t:
        return "KOTA CILEGON"
    if "PANDEGLANG" in t:
        return "KAB. PANDEGLANG"
    if "LEBAK" in t:
        return "KAB. LEBAK"
    return ""


def extract_ruas(letak: str, ket: str) -> str:
    blob = f"{letak}\n{ket}"
    m = re.search(r"Ruas Jalan[^\n,;()]{6,80}", blob, re.I)
    if m:
        name = re.sub(r"\s+", " ", m.group(0)).strip(" -")
        name = re.split(r"\s+[-–]\s+(?:Lampu|PJU|Belanja)", name, maxsplit=1, flags=re.I)[0].strip()
        if len(name) >= 10:
            return name[:90]
    L = first_line(letak)
    if L and re.search(r"\s[-–]\s", L) and not PRODUCT_RE.search(L) and not PROCUREMENT_RE.search(L):
        return L[:90]
    if L and re.match(r"(jl\.?|jalan)\s", L, re.I) and not PRODUCT_RE.search(L) and not JUNK_NAME.search(L):
        return L[:90]
    m = re.search(r"pada Ruas Jalan ([^\n,;()]+)", blob, re.I)
    if m:
        return ("Ruas Jalan " + m.group(1).strip())[:90]
    m = re.search(r"wilayah[^\n(]*\(([^)]{6,80})\)", blob, re.I)
    if m:
        return m.group(1).split(",")[0].strip()[:90]
    return ""


def extract_schools(text: str) -> list[str]:
    """Ambil nama sekolah/yayasan dari kurung pada Merk/Keterangan."""
    t = clean(text).replace("`", "'")
    groups = re.findall(r"\(([^)]{4,100})\)", t)
    if t.count("(") > t.count(")"):
        tail = re.search(r"\(([^)]{4,100})$", t)
        if tail:
            groups.append(tail.group(1))
    out: list[str] = []
    school_tok = re.compile(
        r"\b(SDN|SDS|SMPN|SMP|SMAN|SMA|SMK|MIN|MI|MTs|MAN|Yayasan|Sekolah)\b",
        re.I,
    )
    for g in groups:
        if not school_tok.search(g):
            continue
        chunks = re.split(
            r",\s*(?=(?:SDN|SDS|SD |SMPN|SMP |SMAN|SMA |SMK|MIN |MI |MTs|MAN |MA |Yayasan))",
            g,
            flags=re.I,
        )
        for c in chunks:
            c = re.sub(r"\s+", " ", c).strip(" ,")
            if len(c) < 5:
                continue
            if PROCUREMENT_RE.search(c) or JUNK_NAME.search(c):
                continue
            out.append(c[:90])
    if out:
        return out
    m = re.search(r"ZOSS[^\n]{0,40}Wilayah\s+((?:Kota|Kabupaten|Kab\.)[^\n/,-]+)", t, re.I)
    if m:
        wilayah = re.split(r"\s+Warning|\s+\d+\s+unit", m.group(1).strip())[0].strip()
        if wilayah:
            return [f"Zona Selamat Sekolah {wilayah}"[:90]]
    return []


def extract_d_item(klas: str, letak: str, ket: str) -> tuple[str, str | None]:
    """Nama barang + tipe dari Letak/Ket, bukan klasifikasi KIB."""
    L = first_line(letak)
    K = first_line(ket)
    if L and JUNK_NAME.search(L):
        L = ""
    if K and JUNK_NAME.search(K):
        K = ""
    m = re.search(
        r"(lampu pju[^/\n]{0,50}|pju single|apj/?\s*lpju[^/\n-]{0,40}|"
        r"lpju solar[^/\n]{0,40}|armature[^/\n]{0,50}|stang ornamen|"
        r"tiang (?:oktagonal|double arm|lampu)[^/\n,]{0,50}|"
        r"siklon[^/\n]{0,40})",
        f"{L}\n{K}\n{ket}",
        re.I,
    )
    found = re.sub(r"\s+", " ", m.group(0)).strip(" -") if m else ""
    if found and JUNK_NAME.search(found):
        found = ""
    if L and PRODUCT_RE.search(L) and not PROCUREMENT_RE.search(L):
        return L[:160], klas
    if found:
        return found[:160], (L if L and PRODUCT_RE.search(L) else klas)
    if K and PRODUCT_RE.search(K) and not PROCUREMENT_RE.search(K):
        return K[:160], L or klas
    if L and PROCUREMENT_RE.search(L):
        m2 = re.search(r"(LPJU/?APJ|APJ/?LPJU|PJU|Armature|Lampu)[^/\n]{0,50}", L, re.I)
        return (m2.group(0).strip()[:160] if m2 else klas), klas
    if L and not PROCUREMENT_RE.search(L) and not re.search(r"\s[-–]\s", L):
        return L[:160], klas
    return klas, (L or None)


def split_merk(raw: str) -> tuple[str | None, str | None]:
    s = first_line(raw)
    if not s or s in {"-", "–"}:
        return None, None
    if "/" in s:
        a, b = [p.strip() for p in s.split("/", 1)]
        return (a or None), (b or None)
    return s, None


def site_key(name: str) -> str:
    s = re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-")
    return s[:36] or "lokasi"


def cat_topo_b(name: str, kode: str, ket: str) -> tuple[str, str, str, str | None]:
    n = f"{name} {ket}"
    if VEHICLE_B.search(name) or kode.startswith("1.3.2.02"):
        return "kendaraan", "indoor", "sekretariat", "room-parkir"
    if OUTDOOR_B.search(n) or kode.startswith("1.3.2.04"):
        bidang = "lalin" if re.search(r"rambu|traffic", n, re.I) else "prasjal"
        return "perlengkapan_jalan", "outdoor", bidang, None
    if re.search(r"\bjam\b|absensi|fingerprint", name, re.I):
        return "jam", "indoor", "sekretariat", "room-lobby"
    if "UPTD" in ket:
        return "alat_kantor", "indoor", "upt-terminal", "room-gudang-uptd"
    if re.search(r"meja|kursi|lemari|rak ", name, re.I):
        return "alat_kantor", "indoor", "sekretariat", "room-gudang-sek"
    return "alat_kantor", "indoor", "sekretariat", "room-gudang-atk"


def cat_topo_c(name: str, alamat: str, ket: str) -> tuple[str, str, str, str | None]:
    n = f"{name} {ket} {alamat}".lower()
    if "perlintasan" in n:
        return "gedung", "outdoor", "lautudara", None
    if "pos jaga" in n:
        return "gedung", "outdoor", "lautudara", None
    if any(k in n for k in ("halte", "shelter")):
        return "gedung", "outdoor", "angkutan", None
    if any(k in n for k in ("stasiun", "terminal", "pelabuhan", "menara suar")):
        return "gedung", "outdoor", "lautudara", None
    if "olah raga" in n or "voli" in n:
        return "gedung", "indoor", "sekretariat", "room-voli"
    if "parkir" in n or "garasi" in n or "auning" in n:
        return "gedung", "indoor", "sekretariat", "room-parkir"
    if "ibadah" in n or "mushola" in n:
        return "gedung", "indoor", "sekretariat", "room-dwp"
    if "syeh nawawi" in n or "syech nawawi" in n or "kp3b" in n:
        return "gedung", "indoor", "sekretariat", "room-gudang-atk"
    return "gedung", "outdoor", "sekretariat", None


def iter_rows(path: Path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    for row in ws.iter_rows(min_row=14, values_only=True):
        yield row
    wb.close()


def main() -> None:
    sites: dict[str, dict] = {}
    assets: list[dict] = []
    seen_reg: set[str] = set()

    def add_site(label: str, kab: str, extra: str, site_type: str) -> str | None:
        name = first_line(label) or first_line(extra)
        if not name or PROCUREMENT_RE.search(name) or PRODUCT_RE.search(name) and "ruas" not in name.lower():
            # still allow "Ruas Jalan ..." even if contains PJU later
            if not name or not name.lower().startswith("ruas"):
                name = extract_ruas(label, extra)
        if not name:
            return None
        name = re.sub(r"\s+", " ", name)[:90]
        key = site_key(f"{name}-{kab}")
        sid = f"s-{key}"
        if sid not in sites:
            sites[sid] = {
                "id": sid,
                "code": f"LOC.{len(sites)+1:04d}",
                "name": name,
                "corridor": name[:80],
                "kabupaten": kab or None,
                "lat": None,
                "lng": None,
                "site_type": site_type,
                "desc": (extra[:180] if extra else None),
                "sort": len(sites) + 1,
            }
        return sid

    def put_fixed_site(sid: str, name: str, site_type: str) -> str:
        if sid not in sites:
            sites[sid] = {
                "id": sid,
                "code": sid.replace("site-", "OUT-").upper()[:12],
                "name": name,
                "corridor": name,
                "kabupaten": "Provinsi Banten",
                "lat": None,
                "lng": None,
                "site_type": site_type,
                "desc": name,
                "sort": 200 + len(sites),
            }
        return sid

    put_fixed_site("site-pos-pelintasan", "Pos Pelintasan Kereta Api", "pos")
    put_fixed_site("site-pos-jaga", "Pos Jaga Permanen Jalan", "pos")

    # ----- KIB B: Nama Barang, Merk/Type, Ukuran, Tahun Perolehan, Keterangan -----
    for row in iter_rows(ATT / "KIB B PERALATAN DAN MESIN.xlsx"):
        if not is_main_no(row[0]):
            continue
        kode, idb = parse_ids(row[1])
        if not kode.startswith("1."):
            continue
        name = first_line(row[3])
        if not name or name.lower() in {"perolehan", "pemeliharaan"}:
            continue
        ket = clean(row[17] if len(row) > 17 else "")
        brand, tipe = split_merk(row[4] if len(row) > 4 else "")
        ukuran = first_line(row[5] if len(row) > 5 else "")
        if ukuran in {"-", "/", "-/"}:
            ukuran = ""
        spec_bits = [x for x in [tipe, ukuran] if x]
        spec = " · ".join(spec_bits) or None
        material = first_line(row[6] if len(row) > 6 else "") or None
        year = parse_year(row[7])
        nopol = format_plate(first_line(row[11] if len(row) > 11 else ""))
        bpkb = dash_none(first_line(row[12] if len(row) > 12 else ""))
        rangka = dash_none(first_line(row[9] if len(row) > 9 else ""))
        cat, topo, bidang, room = cat_topo_b(name, kode, ket)
        serial = nopol if nopol else (None if cat == "kendaraan" else (rangka or None))
        note_bits = []
        if cat == "kendaraan" and bpkb:
            note_bits.append(f"BPKB {bpkb}")
        if ket:
            note_bits.append(ket[:240])
        notes = " · ".join(note_bits) or None
        raw_merk = first_line(row[4] if len(row) > 4 else "")
        schools = extract_schools(f"{raw_merk}\n{ket}\n{name}")
        if topo == "outdoor" and schools and raw_merk and "(" in raw_merk:
            head = raw_merk.split("(")[0].strip(" +")
            if head:
                name = head[:160]
                brand, tipe = None, None
                spec = None
        site = lat = lng = None
        if topo == "outdoor":
            loc = ""
            stype = "ruas"
            if schools:
                try:
                    pick = int(idb or 0) % len(schools)
                except ValueError:
                    pick = 0
                loc = schools[pick]
                stype = "sekolah"
            loc = loc or extract_ruas(ket, name)
            kab = kab_from(f"{loc}\n{ket}\n{raw_merk}")
            site = add_site(loc or "", kab, ket or raw_merk, stype)
            if site:
                lat, lng = sites[site]["lat"], sites[site]["lng"]
        reg = f"B-{idb or parse_reg(row[2]) or row[0]}"
        if reg in seen_reg:
            reg = f"{reg}-{row[0]}"
        seen_reg.add(reg)
        assets.append(
            dict(
                id=f"kib-b-{idb or row[0]}",
                register=reg,
                kib_code=kode[:40],
                kib_group="B",
                category=cat,
                topology=topo,
                name=name[:160],
                spec=spec,
                brand=brand,
                material=material,
                year=year,
                qty=parse_qty(row[14] if len(row) > 14 else 1),
                unit="unit",
                condition=parse_cond(row[13] if len(row) > 13 else "", ket),
                room=room,
                site=site,
                bidang=bidang,
                source="APBD",
                serial=serial,
                notes=notes,
                lat=lat,
                lng=lng,
            )
        )

    # ----- KIB C: Nama, Tahun, Letak/Alamat, Ket -----
    for row in iter_rows(ATT / "KIB C GEDUNG DAN BANGUNAN.xlsx"):
        if not is_main_no(row[0]):
            continue
        kode, idb = parse_ids(row[1])
        if not kode.startswith("1."):
            continue
        name = first_line(row[3])
        if not name:
            continue
        alamat = clean(row[9] if len(row) > 9 else "")
        ket = clean(row[16] if len(row) > 16 else "")
        year = parse_year(row[4])
        cond = parse_cond(row[5] if len(row) > 5 else "", ket)
        luas = first_line(row[8] if len(row) > 8 else "")
        cat, topo, bidang, room = cat_topo_c(name, alamat, ket)
        site = lat = lng = None
        blob = f"{alamat}\n{ket}".lower()
        if "perlintasan" in blob:
            site = "site-pos-pelintasan"
            lat, lng = sites[site]["lat"], sites[site]["lng"]
            topo, bidang, room = "outdoor", "lautudara", None
        elif "pos jaga" in blob or "pos jaga" in name.lower():
            site = "site-pos-jaga"
            lat, lng = sites[site]["lat"], sites[site]["lng"]
            topo, bidang, room = "outdoor", "lautudara", None
        elif topo == "outdoor":
            loc = first_line(alamat) or extract_ruas(alamat, ket)
            kab = kab_from(alamat + "\n" + ket)
            stype = "halte" if "halte" in name.lower() else ("pelabuhan" if "pelabuhan" in name.lower() else "area")
            site = add_site(loc or "", kab, ket or alamat, stype)
            if site:
                lat, lng = sites[site]["lat"], sites[site]["lng"]
        spec = ", ".join(x for x in [luas and f"Luas {luas} m2", year and f"TA {year}"] if x) or None
        reg = f"C-{idb or parse_reg(row[2]) or row[0]}"
        if reg in seen_reg:
            reg = f"{reg}-{row[0]}"
        seen_reg.add(reg)
        assets.append(
            dict(
                id=f"kib-c-{idb or row[0]}",
                register=reg,
                kib_code=kode[:40],
                kib_group="C",
                category=cat,
                topology=topo,
                name=name[:160],
                spec=spec,
                brand=None,
                material=first_line(row[7] if len(row) > 7 else "") or None,
                year=year,
                qty=1,
                unit="unit",
                condition=cond,
                room=room,
                site=site,
                bidang=bidang,
                source="APBD",
                serial=None,
                notes=(ket[:280] or alamat[:280] or None),
                lat=lat,
                lng=lng,
            )
        )

    # ----- KIB D: klasifikasi, Letak = nama/tipe atau ruas, Ket, Tahun -----
    for row in iter_rows(ATT / "KIB D JALAN, IRIGASI, DAN JARINGAN.xlsx"):
        if not is_main_no(row[0]):
            continue
        kode, idb = parse_ids(row[1])
        if not kode.startswith("1."):
            continue
        klas = first_line(row[3])
        if not klas:
            continue
        letak = clean(row[9] if len(row) > 9 else "")
        ket = clean(row[18] if len(row) > 18 else "")
        if (not first_line(letak)) and re.search(r"ID AWAL", ket, re.I):
            continue
        year = parse_year(row[4])
        name, tipe = extract_d_item(klas, letak, ket)
        loc = extract_ruas(letak, ket)
        if not loc:
            L = first_line(letak)
            if L and re.search(r"\s[-–]\s", L) and not PRODUCT_RE.search(L):
                loc = L
        kab = kab_from(letak + "\n" + ket)
        site = add_site(loc or "", kab, ket or letak, "ruas")
        lat = lng = None
        if site:
            lat, lng = sites[site]["lat"], sites[site]["lng"]
        spec = tipe if tipe and tipe.lower() != name.lower() else (klas if klas.lower() != name.lower() else None)
        if year and spec:
            spec = f"{spec} · TA {year}"
        elif year:
            spec = f"TA {year}"
        reg = f"D-{idb or parse_reg(row[2]) or row[0]}"
        if reg in seen_reg:
            reg = f"{reg}-{row[0]}"
        seen_reg.add(reg)
        assets.append(
            dict(
                id=f"kib-d-{idb or row[0]}",
                register=reg,
                kib_code=kode[:40],
                kib_group="D",
                category="jaringan",
                topology="outdoor",
                name=name[:160],
                spec=spec[:160] if spec else None,
                brand=None,
                material=first_line(row[5] if len(row) > 5 else "") or None,
                year=year,
                qty=1,
                unit="titik",
                condition=parse_cond(row[15] if len(row) > 15 else "", row[14] if len(row) > 14 else "", ket),
                room=None,
                site=site,
                bidang="prasjal",
                source="APBD",
                serial=None,
                notes=(ket[:280] or None),
                lat=lat,
                lng=lng,
            )
        )

    seen_id: set[str] = set()
    uniq: list[dict] = []
    for a in assets:
        if a["id"] in seen_id:
            a["id"] = f"{a['id']}-x"
        seen_id.add(a["id"])
        uniq.append(a)
    assets = uniq

    payload = {
        "sites": [
            [
                s["id"],
                s["code"],
                s["name"],
                s["corridor"],
                s["kabupaten"],
                s["lat"],
                s["lng"],
                s["site_type"],
                s["desc"],
                s["sort"],
            ]
            for s in sites.values()
        ],
        "assets": [
            [
                a["id"],
                a["register"],
                a["kib_code"],
                a["kib_group"],
                a["category"],
                a["topology"],
                a["name"],
                a["spec"],
                a["brand"],
                a["material"],
                a["year"],
                a["qty"],
                a["unit"],
                a["condition"],
                a["room"],
                a["site"],
                a["bidang"],
                a["source"],
                a["serial"],
                a["notes"],
                a["lat"],
                a["lng"],
            ]
            for a in assets
        ],
    }
    data_dir = ROOT / "data"
    data_dir.mkdir(exist_ok=True)
    raw = json.dumps(payload, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
    gz_path = data_dir / "kib-seed.json.gz"
    gz_path.write_bytes(gzip.compress(raw, compresslevel=9))
    named = sum(1 for a in assets if a["name"] and not a["name"].startswith("Jaringan"))
    geo = sum(1 for a in assets if a["lat"] is not None)
    print("sites", len(sites), "assets", len(assets), "named_goods", named, "with_geo", geo, "kb", round(gz_path.stat().st_size / 1024, 1))
    shown = 0
    for a in assets:
        if a["kib_group"] == "D" and shown < 10:
            shown += 1
            print(f"  {a['year']} | {a['name'][:52]:52} | {(a['spec'] or '')[:28]:28} | {a['site']}")


if __name__ == "__main__":
    main()
